import os, io, csv
from datetime import datetime, date
from functools import wraps
from urllib.parse import urlparse
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file, abort
from flask_wtf.csrf import CSRFProtect
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

app = Flask(__name__)
CSRFProtect(app)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'CHANGE-ME-before-deployment')
database_url = os.environ.get('DATABASE_URL', 'sqlite:///metervault.db')
if database_url.startswith('postgres://'):
    database_url = database_url.replace('postgres://', 'postgresql://', 1)
app.config['SQLALCHEMY_DATABASE_URI'] = database_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['SESSION_COOKIE_SECURE'] = os.environ.get('COOKIE_SECURE', '0') == '1'
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
db = SQLAlchemy(app)

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(20), nullable=False, default='user')
    active = db.Column(db.Boolean, nullable=False, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)

    def set_password(self, password): self.password_hash = generate_password_hash(password)
    def check_password(self, password): return check_password_hash(self.password_hash, password)

class Center(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(160), nullable=False)
    active = db.Column(db.Boolean, nullable=False, default=True)
    sort_order = db.Column(db.Integer, nullable=False, default=0)
    categories = db.relationship('Category', backref='center', cascade='all, delete-orphan', order_by='Category.sort_order')

class Category(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    center_id = db.Column(db.Integer, db.ForeignKey('center.id'), nullable=False)
    name = db.Column(db.String(160), nullable=False)
    active = db.Column(db.Boolean, nullable=False, default=True)
    sort_order = db.Column(db.Integer, nullable=False, default=0)
    meters = db.relationship('Meter', backref='category', cascade='all, delete-orphan', order_by='Meter.sort_order')

class Meter(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    category_id = db.Column(db.Integer, db.ForeignKey('category.id'), nullable=False)
    name = db.Column(db.String(160), nullable=False)
    meter_code = db.Column(db.String(100), nullable=True)
    resource_type = db.Column(db.String(40), nullable=False) # Elekter, Külm vesi, Soe vesi, Gaas, Muu
    unit = db.Column(db.String(20), nullable=False, default='m³')
    active = db.Column(db.Boolean, nullable=False, default=True)
    sort_order = db.Column(db.Integer, nullable=False, default=0)
    readings = db.relationship('Reading', backref='meter', cascade='all, delete-orphan')

class Reading(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    meter_id = db.Column(db.Integer, db.ForeignKey('meter.id'), nullable=False)
    period = db.Column(db.String(7), nullable=False) # YYYY-MM
    value = db.Column(db.Float, nullable=False)
    note = db.Column(db.String(500), nullable=True)
    entered_by = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    entered_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow, onupdate=datetime.utcnow)
    __table_args__ = (db.UniqueConstraint('meter_id', 'period', name='uq_meter_period'),)
    user = db.relationship('User')

class AuditLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    action = db.Column(db.String(100), nullable=False)
    details = db.Column(db.String(500), nullable=False, default='')
    created_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    user = db.relationship('User')

def current_user():
    uid = session.get('user_id')
    return db.session.get(User, uid) if uid else None

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        user = current_user()
        if not user or not user.active:
            session.clear()
            return redirect(url_for('login'))
        return fn(*args, **kwargs)
    return wrapped

def admin_required(fn):
    @wraps(fn)
    @login_required
    def wrapped(*args, **kwargs):
        if current_user().role != 'admin': abort(403)
        return fn(*args, **kwargs)
    return wrapped

def log_action(action, details=''):
    db.session.add(AuditLog(user_id=current_user().id, action=action, details=details[:500]))

def valid_period(p):
    try:
        datetime.strptime(p, '%Y-%m')
        return len(p) == 7
    except (ValueError, TypeError): return False

def previous_reading(meter_id, period):
    r = Reading.query.filter(Reading.meter_id == meter_id, Reading.period < period).order_by(Reading.period.desc()).first()
    return r

@app.template_filter('num')
def format_number(value):
    if value is None: return ''
    try:
        return ('{:,.6f}'.format(float(value)).rstrip('0').rstrip('.')).replace(',', ' ')
    except (TypeError, ValueError):
        return str(value)

@app.context_processor
def inject_globals():
    return {'logged_in_user': current_user(), 'current_year': date.today().year}

@app.route('/')
def index():
    if not current_user(): return redirect(url_for('login'))
    period = request.args.get('period', date.today().strftime('%Y-%m'))
    if not valid_period(period): period = date.today().strftime('%Y-%m')
    centers = Center.query.filter_by(active=True).order_by(Center.sort_order, Center.name).all()
    stats = {}
    for c in centers:
        meters = [m for cat in c.categories if cat.active for m in cat.meters if m.active]
        done = Reading.query.filter(Reading.period == period, Reading.meter_id.in_([m.id for m in meters])).count() if meters else 0
        stats[c.id] = (done, len(meters))
    return render_template('index.html', centers=centers, stats=stats, period=period)

@app.route('/login', methods=['GET','POST'])
def login():
    if current_user(): return redirect(url_for('index'))
    if request.method == 'POST':
        username = request.form.get('username','').strip()
        password = request.form.get('password','')
        user = User.query.filter_by(username=username, active=True).first()
        if user and user.check_password(password):
            session.clear(); session['user_id'] = user.id
            return redirect(url_for('index'))
        flash('Vale kasutajanimi või parool.', 'error')
    return render_template('login.html')

@app.route('/logout', methods=['POST'])
@login_required
def logout():
    session.clear(); return redirect(url_for('login'))

@app.route('/center/<int:center_id>')
@login_required
def center_view(center_id):
    center = db.get_or_404(Center, center_id)
    period = request.args.get('period', date.today().strftime('%Y-%m'))
    if not valid_period(period): period = date.today().strftime('%Y-%m')
    categories = [c for c in center.categories if c.active]
    rows = []
    total = completed = 0
    for cat in categories:
        meters = [m for m in cat.meters if m.active]
        for meter in meters:
            total += 1
            reading = Reading.query.filter_by(meter_id=meter.id, period=period).first()
            if reading: completed += 1
            rows.append({'meter': meter, 'category': cat, 'reading': reading, 'previous': previous_reading(meter.id, period)})
    return render_template('center.html', center=center, categories=categories, rows=rows, period=period, total=total, completed=completed)

@app.route('/category/<int:category_id>/readings', methods=['GET','POST'])
@login_required
def category_readings(category_id):
    category = db.get_or_404(Category, category_id)
    period = request.values.get('period', date.today().strftime('%Y-%m'))
    if not valid_period(period): period = date.today().strftime('%Y-%m')
    meters = [m for m in category.meters if m.active]
    if request.method == 'POST':
        saved = 0
        errors = []
        for meter in meters:
            raw = request.form.get(f'value_{meter.id}', '').strip().replace(',', '.')
            note = request.form.get(f'note_{meter.id}', '').strip()
            if raw == '': continue # blank means leave unchanged, not zero
            try:
                value = float(raw)
                if value < 0: raise ValueError()
            except ValueError:
                errors.append(f'{meter.name}: sisesta nullist suurem või võrdne arv.'); continue
            previous = previous_reading(meter.id, period)
            if previous and value < previous.value:
                errors.append(f'{meter.name}: uus näit on väiksem kui eelmine ({previous.value:g}). Kontrolli väärtust.')
                continue
            reading = Reading.query.filter_by(meter_id=meter.id, period=period).first()
            if reading:
                reading.value = value; reading.note = note; reading.entered_by = current_user().id; reading.updated_at = datetime.utcnow()
            else:
                reading = Reading(meter_id=meter.id, period=period, value=value, note=note, entered_by=current_user().id)
                db.session.add(reading)
            saved += 1
        if saved:
            log_action('readings_saved', f'{category.center.name} / {category.name} / {period}: {saved} näitu')
            db.session.commit()
            flash(f'Salvestatud näite: {saved}.', 'success')
        if errors:
            for error in errors: flash(error, 'error')
        return redirect(url_for('category_readings', category_id=category.id, period=period))
    rows = []
    for m in meters:
        rows.append({'meter': m, 'reading': Reading.query.filter_by(meter_id=m.id, period=period).first(), 'previous': previous_reading(m.id, period)})
    return render_template('category.html', category=category, rows=rows, period=period)

@app.route('/history')
@login_required
def history():
    period = request.args.get('period', date.today().strftime('%Y-%m'))
    if not valid_period(period): period = date.today().strftime('%Y-%m')
    center_id = request.args.get('center_id', type=int)
    query = Reading.query.join(Meter).join(Category).join(Center).filter(Reading.period == period)
    if center_id: query = query.filter(Center.id == center_id)
    readings = query.order_by(Center.sort_order, Center.name, Category.sort_order, Meter.sort_order).all()
    centers = Center.query.order_by(Center.sort_order, Center.name).all()
    return render_template('history.html', readings=readings, centers=centers, period=period, center_id=center_id)

@app.route('/history/<int:reading_id>/edit', methods=['POST'])
@admin_required
def edit_reading(reading_id):
    reading = db.get_or_404(Reading, reading_id)
    raw_value = request.form.get('value', '').strip().replace(',', '.')
    note = request.form.get('note', '').strip()

    try:
        new_value = float(raw_value)
        if new_value < 0 or not __import__('math').isfinite(new_value):
            raise ValueError()
    except ValueError:
        flash('Sisesta korrektne nullist suurem või võrdne näit.', 'error')
        return redirect(url_for(
            'history',
            period=reading.period,
            center_id=reading.meter.category.center_id
        ))

    old_value = reading.value
    old_note = reading.note or ''
    reading.value = new_value
    reading.note = note or None
    reading.updated_at = datetime.utcnow()

    log_action(
        'reading_edited',
        f'{reading.meter.category.center.name} / '
        f'{reading.meter.category.name} / {reading.meter.name}; '
        f'period={reading.period}; value={old_value} -> {new_value}; '
        f'note_changed={old_note != (note or "")}'
    )
    db.session.commit()
    flash('Näidu muudatus salvestatud.', 'success')
    return redirect(url_for(
        'history',
        period=reading.period,
        center_id=reading.meter.category.center_id
    ))


@app.route('/export.xlsx')
@login_required
def export_xlsx():
    period = request.args.get('period', date.today().strftime('%Y-%m'))
    if not valid_period(period): period = date.today().strftime('%Y-%m')
    center_id = request.args.get('center_id', type=int)
    query = Reading.query.join(Meter).join(Category).join(Center).filter(Reading.period == period)
    if center_id: query = query.filter(Center.id == center_id)
    readings = query.order_by(Center.sort_order, Center.name, Category.sort_order, Meter.sort_order).all()
    wb = Workbook(); ws = wb.active; ws.title = 'Näidud'
    ws.append(['Periood','Kaubanduskeskus','Kategooria','Mõõdik','Mõõdiku kood','Ressurss','Ühik','Eelmine näit','Praegune näit','Tarbitud','Märkus','Sisestaja','Sisestatud'])
    for r in readings:
        prev = previous_reading(r.meter_id, period)
        consumption = r.value - prev.value if prev else None
        ws.append([period, r.meter.category.center.name, r.meter.category.name, r.meter.name, r.meter.meter_code or '', r.meter.resource_type, r.meter.unit, prev.value if prev else None, r.value, consumption, r.note or '', r.user.username if r.user else '', r.entered_at.strftime('%Y-%m-%d %H:%M')])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='174A6E'); cell.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = min(max(max(len(str(c.value or '')) for c in col) + 2, 12), 32)
        ws.column_dimensions[get_column_letter(col[0].column)].width = width
    output = io.BytesIO(); wb.save(output); output.seek(0)
    log_action('excel_export', f'period={period}; center={center_id or "all"}'); db.session.commit()
    filename = f'metervault_naidud_{period}' + (f'_center_{center_id}' if center_id else '') + '.xlsx'
    return send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# Admin panel
@app.route('/admin')
@admin_required
def admin_home():
    centers = Center.query.order_by(Center.sort_order, Center.name).all()
    users = User.query.order_by(User.username).all()
    logs = AuditLog.query.order_by(AuditLog.created_at.desc()).limit(30).all()
    return render_template('admin.html', centers=centers, users=users, logs=logs)

@app.route('/admin/center', methods=['POST'])
@admin_required
def admin_add_center():
    name = request.form.get('name','').strip()
    if not name: flash('Sisesta keskuse nimi.', 'error')
    else:
        db.session.add(Center(name=name, sort_order=request.form.get('sort_order', 0, type=int)))
        log_action('center_created', name); db.session.commit(); flash('Kaubanduskeskus lisatud.', 'success')
    return redirect(url_for('admin_home'))

@app.route('/admin/center/<int:center_id>/edit', methods=['POST'])
@admin_required
def admin_edit_center(center_id):
    c = db.get_or_404(Center, center_id); name = request.form.get('name','').strip()
    if name: c.name = name
    c.sort_order = request.form.get('sort_order', c.sort_order, type=int)
    log_action('center_updated', c.name); db.session.commit(); flash('Muudatused salvestatud.', 'success')
    return redirect(url_for('admin_home'))

@app.route('/admin/center/<int:center_id>/archive', methods=['POST'])
@admin_required
def admin_archive_center(center_id):
    c = db.get_or_404(Center, center_id); c.active = not c.active
    log_action('center_active_changed', f'{c.name}: active={c.active}'); db.session.commit()
    flash('Keskuse olek muudetud.', 'success'); return redirect(url_for('admin_home'))

@app.route('/admin/category', methods=['POST'])
@admin_required
def admin_add_category():
    name = request.form.get('name','').strip(); center_id = request.form.get('center_id', type=int)
    center = db.session.get(Center, center_id) if center_id else None
    if not name or not center: flash('Vali kaubanduskeskus ja sisesta kategooria nimi.', 'error')
    else:
        db.session.add(Category(center_id=center.id, name=name, sort_order=request.form.get('sort_order', 0, type=int)))
        log_action('category_created', f'{center.name} / {name}'); db.session.commit(); flash('Kategooria lisatud.', 'success')
    return redirect(url_for('admin_home'))

@app.route('/admin/category/<int:category_id>/edit', methods=['POST'])
@admin_required
def admin_edit_category(category_id):
    c = db.get_or_404(Category, category_id); name = request.form.get('name','').strip()
    if name: c.name = name
    c.sort_order = request.form.get('sort_order', c.sort_order, type=int)
    log_action('category_updated', c.name); db.session.commit(); flash('Kategooria uuendatud.', 'success'); return redirect(url_for('admin_home'))

@app.route('/admin/category/<int:category_id>/archive', methods=['POST'])
@admin_required
def admin_archive_category(category_id):
    c = db.get_or_404(Category, category_id); c.active = not c.active
    log_action('category_active_changed', f'{c.name}: active={c.active}'); db.session.commit(); flash('Kategooria olek muudetud.', 'success'); return redirect(url_for('admin_home'))

@app.route('/admin/meter', methods=['POST'])
@admin_required
def admin_add_meter():
    category_id = request.form.get('category_id', type=int); category = db.session.get(Category, category_id) if category_id else None
    name = request.form.get('name','').strip(); resource = request.form.get('resource_type','Muu'); unit = request.form.get('unit','').strip()
    allowed = {'Elekter','Külm vesi','Soe vesi','Gaas','Muu'}
    if not category or not name or resource not in allowed or not unit:
        flash('Kontrolli kategooriat, mõõdiku nime, ressurssi ja ühikut.', 'error')
    else:
        m = Meter(category_id=category.id, name=name, meter_code=request.form.get('meter_code','').strip() or None, resource_type=resource, unit=unit, sort_order=request.form.get('sort_order', 0, type=int))
        db.session.add(m); log_action('meter_created', f'{category.center.name} / {category.name} / {name}'); db.session.commit(); flash('Mõõdik lisatud.', 'success')
    return redirect(url_for('admin_home'))

@app.route('/admin/meter/<int:meter_id>/edit', methods=['POST'])
@admin_required
def admin_edit_meter(meter_id):
    m = db.get_or_404(Meter, meter_id); name = request.form.get('name','').strip(); unit = request.form.get('unit','').strip()
    resource = request.form.get('resource_type','Muu')
    if name and unit and resource in {'Elekter','Külm vesi','Soe vesi','Gaas','Muu'}:
        m.name=name; m.unit=unit; m.resource_type=resource; m.meter_code=request.form.get('meter_code','').strip() or None; m.sort_order=request.form.get('sort_order', m.sort_order, type=int)
        log_action('meter_updated', m.name); db.session.commit(); flash('Mõõdik uuendatud.', 'success')
    else: flash('Kontrolli mõõdiku andmeid.', 'error')
    return redirect(url_for('admin_home'))

@app.route('/admin/meter/<int:meter_id>/archive', methods=['POST'])
@admin_required
def admin_archive_meter(meter_id):
    m = db.get_or_404(Meter, meter_id); m.active = not m.active
    log_action('meter_active_changed', f'{m.name}: active={m.active}'); db.session.commit(); flash('Mõõdiku olek muudetud.', 'success'); return redirect(url_for('admin_home'))

@app.route('/admin/user', methods=['POST'])
@admin_required
def admin_add_user():
    username=request.form.get('username','').strip(); password=request.form.get('password',''); role=request.form.get('role','user')
    if not username or len(password) < 10 or role not in {'user','admin'}:
        flash('Kasutajanimi on kohustuslik ja parool peab olema vähemalt 10 märki.', 'error')
    elif User.query.filter_by(username=username).first(): flash('See kasutajanimi on juba kasutusel.', 'error')
    else:
        u=User(username=username, role=role); u.set_password(password); db.session.add(u); log_action('user_created', f'{username} ({role})'); db.session.commit(); flash('Kasutaja lisatud.', 'success')
    return redirect(url_for('admin_home'))

@app.route('/admin/user/<int:user_id>/toggle', methods=['POST'])
@admin_required
def admin_toggle_user(user_id):
    u=db.get_or_404(User,user_id)
    if u.id == current_user().id: flash('Enda kontot ei saa siin deaktiveerida.', 'error')
    else:
        u.active=not u.active; log_action('user_active_changed', f'{u.username}: active={u.active}'); db.session.commit(); flash('Kasutaja olek muudetud.', 'success')
    return redirect(url_for('admin_home'))

@app.route('/admin/user/<int:user_id>/password', methods=['POST'])
@admin_required
def admin_reset_password(user_id):
    u=db.get_or_404(User,user_id); password=request.form.get('password','')
    if len(password) < 10: flash('Parool peab olema vähemalt 10 märki.', 'error')
    else: u.set_password(password); log_action('password_reset', u.username); db.session.commit(); flash('Parool muudetud.', 'success')
    return redirect(url_for('admin_home'))

@app.errorhandler(403)
def forbidden(_): return render_template('error.html', code=403, message='Sul puudub sellele lehele ligipääs.'), 403
@app.errorhandler(404)
def not_found(_): return render_template('error.html', code=404, message='Lehte ei leitud.'), 404

with app.app_context():
    db.create_all()
    # Optional first-run admin bootstrap. Never hardcode production credentials.
    bootstrap_user = os.environ.get('BOOTSTRAP_ADMIN_USER')
    bootstrap_password = os.environ.get('BOOTSTRAP_ADMIN_PASSWORD')
    if bootstrap_user and bootstrap_password and not User.query.filter_by(role='admin').first():
        admin = User(username=bootstrap_user, role='admin'); admin.set_password(bootstrap_password); db.session.add(admin); db.session.commit()

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)), debug=False)
